#!/usr/bin/env python3
"""Dump a .pptx as plain text so it can be rewritten as a Marp deck.

    pptx-dump.py deck.pptx              # slide texts (indented by level), links, tables, notes
    pptx-dump.py deck.pptx --images DIR # additionally save every picture as DIR/sNN-name.ext

Only the slide *content* is printed; header/footer/slide-number placeholders
of the old template are skipped.
"""
import argparse
import pathlib
import re

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

SKIP_PLACEHOLDERS = {
    PP_PLACEHOLDER.FOOTER,
    PP_PLACEHOLDER.SLIDE_NUMBER,
    PP_PLACEHOLDER.DATE,
    PP_PLACEHOLDER.HEADER,
}


def slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")[:40]


def walk(shapes, slide_no, img_dir, depth=0):
    for sh in shapes:
        ind = "  " * depth
        if sh.is_placeholder and sh.placeholder_format.type in SKIP_PLACEHOLDERS:
            continue
        # The old IW1 template carries the running header in two BODY
        # placeholders glued to the top edge; skip them too.
        if sh.is_placeholder and sh.top == 0 and sh.height <= 300000:
            continue
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            print(f"{ind}[group {sh.name}]")
            walk(sh.shapes, slide_no, img_dir, depth + 1)
            continue
        # Pictures, including those dropped into a content placeholder
        # (python-pptx reports those as PLACEHOLDER, but they carry .image).
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE or hasattr(sh, "image"):
            img = sh.image
            print(f"{ind}[image {img.filename} {img.content_type} {len(img.blob)} B]")
            if img_dir:
                out = img_dir / f"s{slide_no:02d}-{slug(sh.name)}.{img.ext}"
                out.write_bytes(img.blob)
                print(f"{ind}  -> {out}")
            continue
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                print(ind + "| " + " | ".join(c.text.replace("\n", " / ") for c in row.cells) + " |")
            continue
        if sh.has_text_frame:
            tag = "TITLE" if sh.is_placeholder and sh.placeholder_format.type in (
                PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE) else None
            for p in sh.text_frame.paragraphs:
                text = "".join(r.text for r in p.runs).strip()
                if not text:
                    continue
                links = [r.hyperlink.address for r in p.runs if r.hyperlink and r.hyperlink.address]
                links = list(dict.fromkeys(links))
                prefix = f"{ind}# " if tag else f"{ind}{'  ' * p.level}- "
                print(prefix + text + (f"   <{' | '.join(links)}>" if links else ""))


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pptx")
    ap.add_argument("--images", metavar="DIR", help="save pictures into DIR")
    args = ap.parse_args()

    img_dir = pathlib.Path(args.images) if args.images else None
    if img_dir:
        img_dir.mkdir(parents=True, exist_ok=True)

    prs = Presentation(args.pptx)
    print(f"# {args.pptx}: {len(prs.slides)} slides, {prs.slide_width}x{prs.slide_height} EMU")
    for i, slide in enumerate(prs.slides, 1):
        print(f"\n===== slide {i} ({slide.slide_layout.name})")
        walk(slide.shapes, i, img_dir)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                print("NOTES:")
                for line in notes.splitlines():
                    print("  " + line)


if __name__ == "__main__":
    main()
